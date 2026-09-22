#!/usr/bin/env python3
"""Build the editable teacher-facing companion to the canonical content intake.

Maintainers: run with python-docx available, then render and inspect all pages.
The canonical Markdown remains authoritative. This is a concise Word adaptation,
not a second schema. Overview labels, the ownership rule, topic labels, and the
approval fields are read from that source. If its meaning changes, update this
adaptation and regenerate. --check fails when the stored source hash is stale.
Teachers use the Word file directly; they do not run this script.
"""

import argparse
import hashlib
import re
from pathlib import Path

import zipfile

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SPOKES Builder/content-intake-template.md"
OUTPUT = ROOT / "docs/bespoke/team-lesson-intake.docx"
VERSION = "1.0"


def plain(text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").strip()


def canonical_fields(text, section):
    part = text.split(section, 1)[1].split("\n## ", 1)[0]
    return [plain(line.split("|")[1]) for line in part.splitlines()
            if line.startswith("| **")]


def build(text, digest):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.header_distance = sec.footer_distance = Inches(.28)
    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in styles[name]._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Title"].font.size = Pt(25)
    styles["Title"].paragraph_format.space_after = Pt(10)
    styles["Heading 1"].font.size = Pt(19)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 3"].font.size = Pt(11)
    doc.core_properties.title = "SPOKES Team Lesson Intake"
    doc.core_properties.subject = "Editable team planning form for Word in Teams and OneDrive"
    doc.core_properties.author = "SPOKES"
    doc.core_properties.version = VERSION
    doc.core_properties.comments = f"Canonical source: SPOKES Builder/content-intake-template.md; SHA256:{digest}"

    def p(text="", style=None):
        return doc.add_paragraph(text, style)

    def field(label, hint="Type here", after=6):
        para = p()
        para.paragraph_format.space_after = Pt(after)
        para.add_run(label + ": ").bold = True
        para.add_run("[" + hint + "]")
        return para

    def heading(text, level=2):
        doc.add_heading(text, level)

    def page(title):
        doc.add_page_break()
        heading(title, 1)

    def table(headers, widths, rows=2):
        t = doc.add_table(rows=rows+1, cols=len(headers))
        t.autofit = False
        for c, width in zip(t.columns, widths):
            c.width = Inches(width)
        for i, title in enumerate(headers):
            t.rows[0].cells[i].text = title
        for row in t.rows[1:]:
            for cell in row.cells:
                cell.text = "[Type here]"
        for index, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            if index == 0:
                repeat = OxmlElement("w:tblHeader")
                trpr.append(repeat)
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                pr = cell._tc.get_or_add_tcPr()
                margin = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    e = OxmlElement("w:" + edge)
                    e.set(qn("w:w"), "100")
                    e.set(qn("w:type"), "dxa")
                    margin.append(e)
                pr.append(margin)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    e = OxmlElement("w:" + edge)
                    e.set(qn("w:val"), "single")
                    e.set(qn("w:sz"), "4")
                    e.set(qn("w:color"), "D9D9D9")
                    borders.append(e)
                pr.append(borders)
                if index == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "EDF3F7")
                    pr.append(fill)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(3)
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
        p().paragraph_format.space_after = Pt(0)
        return t

    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    footer.add_run(f"SPOKES Team Lesson Intake  |  Template {VERSION}  |  Page ").font.size = Pt(9)
    pagefield = OxmlElement("w:fldSimple")
    pagefield.set(qn("w:instr"), "PAGE")
    footer._p.append(pagefield)

    doc.add_heading("SPOKES Team Lesson Intake", 0)
    p("Use this Word file to agree on the content of one lesson. Type in the brackets, add rows or pages when needed, and keep the shared copy in your team's Teams or OneDrive folder.")
    heading("Start with one shared copy")
    p("Name your copy for the lesson. During a Teams call, one person can share the screen while the team talks through each section. Use comments for suggestions. The team lead records the wording everyone agrees to use.")
    p("Keep working drafts in the same shared folder. Before sending work to the build team, save a clearly named approved copy and confirm the agreed BeSpoke design opens again. OneDrive does not combine separate design files automatically.")
    labels = canonical_fields(text, "## Section 1: Lesson Overview")
    for label in ("Lesson Title", "Lesson Subtitle", "Module Number", "Content Team / Author", "Team Lead / Content Approver", "Date Submitted", "Submission Revision", "Approved BeSpoke Design File", "Target Learners / Delivery"):
        if label not in labels:
            raise ValueError(f"Canonical overview field changed: {label}")
        hint = "Optional" if label == "Module Number" else "Type here"
        if label == "Submission Revision":
            hint = "Version name, such as draft 1 or approved 1"
        if label == "Approved BeSpoke Design File":
            hint = "Filename and saved date; the team lead can fill this in"
        if label == "Target Learners / Delivery":
            hint = "Learners, reading or language supports, length, online or in person"
        field(label, hint)
    heading("Your words stay yours")
    rule = re.search(r"\*\*Instructor ownership:\*\* (.+)", text).group(1)
    p(plain(rule))

    page("1  Begin the lesson")
    p("WIPPEA means Warm Up, Introduction, Presentation, Practice, Evaluation, and Application. Use the prompts below to plan what learners will do. Write N/A for anything that does not apply.")
    field("Lesson description", "Two or three sentences about the lesson and why it matters", 12)
    heading("Warm Up")
    field("Opening activity or reflection", "Exact prompt and simple instructions", 14)
    field("Questions to ask", "Write the questions learners will discuss", 14)
    field("Time and materials", "Minutes, supplies, and handout names", 10)
    heading("Introduction")
    field("Learning objectives", "What will learners be able to do? Keep this exact wording in the lesson", 16)
    field("Why this matters", "Opening explanation, story, or framing statement", 14)
    field("Key terms", "Words learners need and a plain-language definition of each", 14)
    heading("Ways to take part")
    field("Participation choices", "Spoken, written, partner, or other options; no forced personal disclosure", 12)
    field("Support for learners", "Reading, language, disability, technology, or confidence supports")

    page("2  Teach and practice")
    p("Complete one block for each topic. Copy a block when you need more. Identify both the explanation and guided practice, even when they appear in the same chapter. Use the team's exact activity wording.")
    topic_headings = re.findall(r"### Stage (P[123]) -- Presentation ([123])", text)
    if len(topic_headings) != 3:
        raise ValueError("Canonical presentation structure changed; review the Word adaptation")
    for _, n in topic_headings:
        heading(f"Topic {n}")
        field("Topic title", "Type here", 4)
        field("Teaching points", "Facts, framework steps, comparison, or myth and correction", 5)
        field("Guided practice", "What learners do, exact directions, time, and materials", 5)
        field("Questions and feedback", "What to ask, expected reasoning, and help if learners struggle", 5)
        field("Video topic or visual idea", "Describe it if useful; write N/A if none", 5)
    heading("Examples and wording to preserve")
    field("Story or scenario", "Write the exact example or identify the attached file", 9)
    field("Quotation or key statement", "Exact wording, who said it, and the source; add any closing statement")

    page("3  Check learning and apply it")
    heading("Evaluation")
    p("Describe what learners will show or do. A self-rating is useful for reflection; include an observable demonstration of the skill as well.")
    field("Exit ticket or assessment", "Exact questions, task, and instructions", 15)
    field("Expected answers or scoring", "Correct answers with reasons, or clear rubric criteria", 14)
    table(["Objective", "Practice and evidence of learning", "Handout or rubric"], [1.6, 3.7, 1.7], 2)
    field("Equivalent ways to complete it", "Include an offline option and supported participation where needed", 10)
    field("Pre and post test and rubric files", "List the files supplied; note anything still needed", 8)
    heading("Application")
    field("Closing activity", "Exact instructions, discussion questions, materials, and time", 14)
    field("Use outside class", "One action learners can try at home, in class, or at work", 14)
    field("Follow up", "How learners will reflect on or demonstrate what happened", 10)
    field("Closing statement", "Team wording or a verified quotation; record its source in the next section")

    page("4  Keep the resources together")
    p("Keep the lesson PowerPoint, handouts, teacher notes, pre and post test, rubric, and design file in the same team folder. Supply PDF copies for lesson links and keep the editable originals for corrections. Make sure text and meaningful images are accessible.")
    table(["Filename", "What it is and where it is used", "Ready or still needed"], [2.2, 3.3, 1.5], 3)
    heading("Sources for facts and materials")
    p("Record sources for facts, numbers, frameworks, quotations, and borrowed materials. Use current authoritative sources. For workplace rights, benefits, money, health, or safety, record the jurisdiction and a qualified reviewer. Flag uncertainty for the team; agents must not invent a replacement.")
    field("Claim or material and lesson location", "Type here", 7)
    field("Source title author and link or filename", "Type here", 7)
    field("Published or updated date and date checked", "Type here", 7)
    field("Jurisdiction audience and limits", "Where and to whom this information applies", 7)
    field("Permission or license", "Evidence that a borrowed item may be reused, or review still needed", 7)
    p("Copy these source fields for each additional source, or attach the team's source list with the same information. A research link does not give permission to download or redistribute a video.")
    heading("Preferences for the finished lesson")
    field("Tone and visual preferences", "Describe any comparisons, acronyms, images, or activity layouts you want")
    p("Record agreed choices in BeSpoke. The build uses the project's approved branding and accessibility rules. Any necessary change to the agreed design comes back for human review.")

    page("5  Review and hand over")
    p("The team lead uses this page when the team is ready to send its work. An approved intake starts the build process. A teacher still reviews the finished lesson before it is taught.")
    heading("Check together")
    checks = [
        "Every WIPPEA stage has content, including guided practice and application.",
        "Each objective has an activity and an assessment with expected answers or scoring.",
        "Instructor objectives, stories, hooks, quotations, and activity wording remain exact.",
        "Facts and borrowed materials have sources, check dates, and reuse information.",
        "Personal disclosure is optional; accessible and offline alternatives are recorded.",
        "The PowerPoint, handout PDFs and editable originals, teacher notes, test, and rubric are included or explicitly marked N/A with a reason.",
        "The saved BeSpoke design matches this approved intake, and another team member can reopen it.",
        "Any unresolved content or design decision has an owner; nobody will invent missing content.",
        "A named teacher will review wording, resources, assessment, and design in the finished lesson before teaching.",
    ]
    for check in checks:
        p("[  ] " + check)
    heading("Team approval")
    # Approval fields are derived from the canonical table, preserving its scope.
    approval = text.split("### Approval and unresolved decisions", 1)[1]
    fields = [plain(line.split("|")[1]) for line in approval.splitlines()
              if line.startswith("| ") and not line.startswith("| Field")]
    for label in fields:
        field(label, "Type here", 5)
    note = p("Template version " + VERSION + ". Source: SPOKES Builder/content-intake-template.md. This Word form is a concise companion to that template. It grows as your team adds content.")
    for run in note.runs:
        run.font.size = Pt(9)
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Check the generated Word file matches the canonical source revision")
    args = parser.parse_args()
    text = SOURCE.read_text(encoding="utf-8")
    digest = hashlib.sha256(text.encode()).hexdigest()
    if args.check:
        try:
            with zipfile.ZipFile(OUTPUT) as package:
                current = f"SHA256:{digest}".encode() in package.read("docProps/core.xml")
        except (OSError, KeyError, zipfile.BadZipFile):
            current = False
        if not current:
            raise SystemExit("Word intake is missing or stale; review the adaptation, regenerate, and visually check it.")
        print("Word intake source hash matches the canonical template.")
        return
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build(text, digest).save(OUTPUT)
    print(f"Created {OUTPUT.relative_to(ROOT)} from source SHA256 {digest}")


if __name__ == "__main__":
    main()
